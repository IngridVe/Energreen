"""Wekelijkse mediamonitoring: haalt de feeds uit feeds.csv op en schrijft een Excel-overzicht.

Draait in GitHub Actions. Lokaal draaien kan ook:
    pip install feedparser openpyxl
    python scripts/monitor.py --dagen 7
"""

import argparse
import csv
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlparse

import feedparser
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
FEEDS = ROOT / "feeds.csv"
OUTDIR = ROOT / "output"

ARIAL = "Arial"
KOP_FONT = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
KOP_VULLING = PatternFill("solid", fgColor="1F3864")
TEKST = Font(name=ARIAL, size=10)
LINKTEKST = Font(name=ARIAL, size=10, color="0563C1", underline="single")
RAND = Border(bottom=Side(style="thin", color="D9D9D9"))
BE_VULLING = PatternFill("solid", fgColor="E2EFDA")

BE_DOMEINEN = (
    "tijd.be", "hln.be", "standaard.be", "vrt.be", "nieuwsblad.be", "lecho.be",
    "rtbf.be", "sudinfo.be", "lavenir.net", "knack.be", "demorgen.be", "bruzz.be",
    "gva.be", "hbvl.be", "made-in.be", "test-aankoop.be", "test-achats.be",
    "vlaamsenutsregulator.be", "creg.be", "cwape.be", "brugel.brussels",
    "energiesparen.be", "fluvius.be", "elia.be", "belga.be", "trends.knack.be",
)


def lees_feeds():
    with FEEDS.open(encoding="utf-8") as f:
        return [r for r in csv.DictReader(f) if r.get("feed_url", "").strip()]


def naar_datum(entry):
    for veld in ("published_parsed", "updated_parsed"):
        waarde = entry.get(veld)
        if waarde:
            return datetime(*waarde[:6], tzinfo=timezone.utc)
    return None


def bron_van(entry, link):
    bron = entry.get("source", {}).get("title") if isinstance(entry.get("source"), dict) else None
    if bron:
        return bron
    host = urlparse(link).netloc.replace("www.", "")
    return host or "onbekend"


def regio_van(entry, link, titel):
    tekst = " ".join(filter(None, [link, titel, str(entry.get("source", ""))])).lower()
    for domein in BE_DOMEINEN:
        if domein in tekst:
            return "BE"
    return "overig"


def normaliseer(titel):
    return re.sub(r"[^a-z0-9]+", " ", titel.lower()).strip()


def verzamel(dagen):
    grens = datetime.now(timezone.utc) - timedelta(days=dagen)
    rijen, gezien, problemen = [], set(), []

    for feed in lees_feeds():
        geparsed = feedparser.parse(feed["feed_url"])
        if getattr(geparsed, "bozo", False) and not geparsed.entries:
            problemen.append(f"{feed['thema']} ({feed['taal']}): feed niet leesbaar")
            continue
        if not geparsed.entries:
            problemen.append(f"{feed['thema']} ({feed['taal']}): geen items in de feed")
            continue

        for entry in geparsed.entries:
            datum = naar_datum(entry)
            if datum is None or datum < grens:
                continue
            titel = (entry.get("title") or "").strip()
            link = (entry.get("link") or "").strip()
            if not titel or not link:
                continue
            sleutel = normaliseer(titel)
            if sleutel in gezien:
                continue
            gezien.add(sleutel)
            rijen.append({
                "datum": datum.strftime("%Y-%m-%d"),
                "thema": feed["thema"],
                "taal": feed["taal"],
                "regio": regio_van(entry, link, titel),
                "titel": titel,
                "bron": bron_van(entry, link),
                "link": link,
            })

    rijen.sort(key=lambda r: (r["datum"], r["regio"] == "BE"), reverse=True)
    return rijen, problemen


def schrijf_excel(rijen, problemen, dagen, pad):
    wb = Workbook()
    ws = wb.active
    ws.title = "Overzicht"
    koppen = ["Datum", "Thema", "Regio", "Taal", "Titel", "Bron", "Link",
              "Opportuniteit", "Opmerking Energreen"]
    breedtes = [12, 20, 8, 6, 78, 26, 16, 14, 32]
    for i, (kop, breedte) in enumerate(zip(koppen, breedtes), start=1):
        cel = ws.cell(row=1, column=i, value=kop)
        cel.font, cel.fill = KOP_FONT, KOP_VULLING
        ws.column_dimensions[get_column_letter(i)].width = breedte
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(len(rijen) + 1, 2)}"

    for r, rij in enumerate(rijen, start=2):
        be = rij["regio"] == "BE"
        waarden = [rij["datum"], rij["thema"], rij["regio"], rij["taal"], rij["titel"], rij["bron"]]
        for i, waarde in enumerate(waarden, start=1):
            cel = ws.cell(row=r, column=i, value=waarde)
            cel.font = TEKST
            cel.alignment = Alignment(wrap_text=(i == 5), vertical="top")
            cel.border = RAND
            if be:
                cel.fill = BE_VULLING
        cel = ws.cell(row=r, column=7, value="artikel openen")
        cel.hyperlink = rij["link"]
        cel.font = LINKTEKST
        cel.border = RAND
        cel.alignment = Alignment(vertical="top")
        if be:
            cel.fill = BE_VULLING
        for i in (8, 9):
            ws.cell(row=r, column=i).border = RAND
        ws.row_dimensions[r].height = 30

    ws2 = wb.create_sheet("Logboek")
    ws2.column_dimensions["A"].width = 100
    regels = [
        f"Run van {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC",
        f"Venster: de laatste {dagen} dagen",
        f"Aantal items: {len(rijen)}  |  waarvan Belgisch: {sum(1 for r in rijen if r['regio'] == 'BE')}",
        "",
        "Groen in het overzicht = Belgische bron. Links lopen via Google News en openen het artikel bij de titel zelf.",
        "Print, radio en tv zitten niet in deze feeds. Die blijven via Belga lopen.",
        "",
        "Feeds zonder resultaat of met een probleem:",
    ]
    regels += [f"  - {p}" for p in problemen] or ["  - geen"]
    for r, regel in enumerate(regels, start=1):
        cel = ws2.cell(row=r, column=1, value=regel)
        cel.font = Font(name=ARIAL, size=10, bold=(r == 1))
        cel.alignment = Alignment(wrap_text=True, vertical="top")

    pad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pad)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dagen", type=int, default=7)
    args = parser.parse_args()

    rijen, problemen = verzamel(args.dagen)
    datumstempel = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    pad = OUTDIR / f"monitoring_{datumstempel}.xlsx"
    schrijf_excel(rijen, problemen, args.dagen, pad)

    print(f"{len(rijen)} items weggeschreven naar {pad}")
    for p in problemen:
        print(f"  probleem: {p}", file=sys.stderr)


if __name__ == "__main__":
    main()
